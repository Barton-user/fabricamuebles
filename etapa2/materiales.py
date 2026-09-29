# -*- coding: utf-8 -*-
"""
materiales — lista de materiales del paquete de produccion, desde maestra.json.

    python3 materiales.py maestra.json -o <carpeta> [--mueble "00-Bandejero"]

Escribe:
  lista_materiales.xlsx   hojas: Piezas · Herrajes · Resumen
  piezas_autocut.csv      una fila por placa con el codigo LIMPIO (13 digitos) en
                          una columna propia, para mapear directo en AutoCUT
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import sys
from collections import OrderedDict

COL_PIEZAS = ["Mueble", "Globo", "Código", "Pieza", "Cant", "Largo (H)", "Ancho (W)", "Espesor",
              "Corte largo", "Corte ancho", "Material", "Textura", "Veta",
              "Canto izq", "Canto abajo", "Canto der", "Canto arriba", "Rol", "Aviso", "Nombre GuiGui", "Nombre completo GuiGui"]


def _n(v):
    return int(v) if abs(v - round(v)) < 1e-9 else round(v, 2)


def filas_piezas(m, solo=None):
    out = []
    for mu in m["muebles"]:
        if solo and mu["nombre"] not in solo:
            continue
        for p in mu["piezas"]:
            c = p["cantos"]
            out.append([mu["nombre"], p["globo"], p["codigo"], p["nombre"], 1,
                        _n(p["H"]), _n(p["W"]), _n(p["T"]), _n(p["corte_H"]), _n(p["corte_W"]),
                        p["material"], p["textura"], p["veta"],
                        _n(c["izq"]), _n(c["abajo"]), _n(c["der"]), _n(c["arriba"]),
                        p["rol"], p.get("aviso", ""), p.get("nombre_guigui", ""), p["nombre_largo"]])
    return out


def filas_herrajes(m, solo=None):
    out = []
    tot = OrderedDict()

    def suma(k, v):
        tot[k] = tot.get(k, 0) + v
    for mu in m["muebles"]:
        if solo and mu["nombre"] not in solo:
            continue
        h = mu["herrajes"]
        n3 = h["excentricas"]
        filas = [
            ("Tres en uno · excéntrica Ø15", n3, "u"),
            ("Tres en uno · perno Ø8x33", h["pernos"], "u"),
            ("Tres en uno · receptor Ø10", h["receptores"], "u"),
            ("Bisagra cazoleta Ø35 full overlay", h["bisagras"], "u"),
            ("Tornillo bisagra Ø3,5x16 (ala + base)", h["tornillos_bisagra"], "u"),
            ("Corredera telescópica %s (par)" % (h.get("correderas_largo") or "450mm"), h["correderas_pares"], "par"),
            ("Tornillo corredera", h["tornillos_corredera"], "u"),
        ]
        for largo in sorted(set(h["led"])):
            filas.append(("Tira LED %d mm (perfil 9x9)" % largo, h["led"].count(largo), "u"))
        for nombre, cant, unidad in filas:
            if cant:
                out.append([mu["nombre"], nombre, cant, unidad])
                suma((nombre, unidad), cant)
    return out, [["TOTAL", k[0], v, k[1]] for k, v in tot.items()]


def resumen(m, solo=None):
    acc = OrderedDict()
    for mu in m["muebles"]:
        if solo and mu["nombre"] not in solo:
            continue
        for p in mu["piezas"]:
            k = (p["material"], p["textura"], _n(p["T"]))
            a = acc.setdefault(k, [0, 0.0])
            a[0] += 1
            a[1] += p["corte_W"] * p["corte_H"] / 1e6
    return [[k[0], k[1], k[2], v[0], round(v[1], 3)] for k, v in acc.items()]


def escribir(m, salida, solo=None):
    os.makedirs(salida, exist_ok=True)
    piezas = filas_piezas(m, solo)
    herr, tot = filas_herrajes(m, solo)
    res = resumen(m, solo)
    # --- csv para AutoCUT ------------------------------------------------
    with open(os.path.join(salida, "piezas_autocut.csv"), "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(COL_PIEZAS)
        w.writerows(piezas)
    # --- xlsx --------------------------------------------------------------
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("sin openpyxl: solo csv")
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "Piezas"
    ws.append(["%s — orden %s — lista de piezas" % (m.get("ambiente", ""), m.get("orden", ""))])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(COL_PIEZAS)
    for c in ws[2]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDDDDD")
    for r in piezas:
        ws.append(r)
    for i, ancho in enumerate([30, 6, 15, 26, 5, 10, 10, 8, 10, 10, 10, 12, 10, 8, 8, 8, 8, 8, 34, 24, 50], 1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    for row in ws.iter_rows(min_row=3, min_col=3, max_col=3):
        for c in row:
            c.number_format = "@"
            c.alignment = Alignment(horizontal="left")
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = "A2:%s%d" % (get_column_letter(len(COL_PIEZAS)), ws.max_row)

    wh = wb.create_sheet("Herrajes")
    wh.append(["Mueble", "Herraje", "Cantidad", "Unidad"])
    for c in wh[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDDDDD")
    for r in herr:
        wh.append(r)
    wh.append([])
    for r in tot:
        wh.append(r)
        for c in wh[wh.max_row]:
            c.font = Font(bold=True)
    for i, ancho in enumerate([32, 42, 10, 8], 1):
        wh.column_dimensions[get_column_letter(i)].width = ancho

    wr = wb.create_sheet("Resumen")
    wr.append(["Material", "Textura", "Espesor", "Placas", "m² de corte"])
    for c in wr[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDDDDD")
    for r in res:
        wr.append(r)
    wr.append(["TOTAL", "", "", sum(r[3] for r in res), round(sum(r[4] for r in res), 3)])
    for c in wr[wr.max_row]:
        c.font = Font(bold=True)
    for i, ancho in enumerate([14, 14, 9, 8, 12], 1):
        wr.column_dimensions[get_column_letter(i)].width = ancho
    wb.save(os.path.join(salida, "lista_materiales.xlsx"))
    print("%s: %d piezas, %d filas de herrajes, %d materiales"
          % (os.path.join(salida, "lista_materiales.xlsx"), len(piezas), len(herr), len(res)))


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("maestra")
    ap.add_argument("-o", "--salida", default=".")
    ap.add_argument("--mueble", action="append")
    a = ap.parse_args(argv)
    with open(a.maestra, "r", encoding="utf-8") as f:
        m = json.load(f)
    escribir(m, a.salida, a.mueble)
    return 0


if __name__ == "__main__":
    sys.exit(main())
